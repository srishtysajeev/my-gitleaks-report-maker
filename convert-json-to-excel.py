import sys
import pandas as pd
import json
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

# When running the script the name of the json and the excel file must be given in as arguments
json_name = sys.argv[1]
excel_name = sys.argv[2]

print(f"Converting {json_name} to excel")

try:
    with open(json_name) as file:
        data = json.load(file)
    # Need to normalize to get correct headers
    df = pd.json_normalize(data)
    print(df)

except json.decoder.JSONDecodeError:
    print("JSON is empty — creating Excel file with only headers.")
    df = pd.DataFrame() 

# Convert any list-like fields to comma-separated strings
for col in df.columns:
    if df[col].apply(lambda x: isinstance(x, list)).any():
        df[col] = df[col].apply(lambda x: ", ".join(x) if isinstance(x, list) else x)

# Export to Excel
df.to_excel(excel_name, index=False)

# Load workbook
wb = load_workbook(excel_name)
ws = wb.active

# --- AUTO‑DETECT NEXT EMPTY COLUMN ---
# Find the first empty column in row 1
col = 1
while ws.cell(row=1, column=col).value not in (None, ""):
    col += 1

start_col = col  # first free column

# Extra headers to add
extra_headers = [
    "Leak Classification",
    "Impact Assessment",
    "Developer Comments",
    "Mitigation Status",
    "Security Team Comments"
]

# Write headers dynamically
for i, header in enumerate(extra_headers):
    ws.cell(row=1, column=start_col + i, value=header)

# Map header names to their column letters
from openpyxl.utils import get_column_letter
col_letters = {
    header: get_column_letter(start_col + i)
    for i, header in enumerate(extra_headers)
}

# --- DROPDOWN DEFINITIONS ---
dv_class = DataValidation(
    type="list",
    formula1='"Genuine,False Positive,Test,Other"',
    allow_blank=True
)

dv_impact = DataValidation(
    type="list",
    formula1='"Critical,Medium,Low,Informational"',
    allow_blank=True
)

dv_mitig = DataValidation(
    type="list",
    formula1='"Action Taken,Not Applicable,Pending,Completed"',
    allow_blank=True
)

ws.add_data_validation(dv_class)
ws.add_data_validation(dv_impact)
ws.add_data_validation(dv_mitig)

# Apply dropdowns only if rows exist
max_row = ws.max_row
if max_row > 1:
    dv_class.add(f"{col_letters['Leak Classification']}2:{col_letters['Leak Classification']}{max_row}")
    dv_impact.add(f"{col_letters['Impact Assessment']}2:{col_letters['Impact Assessment']}{max_row}")
    dv_mitig.add(f"{col_letters['Mitigation Status']}2:{col_letters['Mitigation Status']}{max_row}")

# Save final file
wb.save(excel_name)
